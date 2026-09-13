import json

import pytest

from xtoo.config import load_settings
from xtoo.extract import extract_text
from xtoo.indexer import Indexer
from xtoo.store import Store


def test_scripts_are_indexed_as_text(library):
    root, _, store, indexer = library
    scripts = {
        "Deploy-Cluster.ps1": "Write-Host 'restart the broker service'",
        "cleanup.bat": "@echo off\nrem purge the broker cache",
        "rotate-logs.sh": "#!/bin/bash\n# archive the broker logs",
        "report.py": "def main():\n    print('broker summary')",
        "schema.sql": "SELECT * FROM broker_events;",
    }
    for name, text in scripts.items():
        (root / name).write_text(text)
    (root / "tool.exe").write_bytes(b"MZ binary payload")
    assert indexer.scan()["indexed"] == len(scripts)
    assert store.search("broker")["total"] == len(scripts)
    assert store.search("Write-Host")["total"] == 1
    assert store.search("restart", "ps1")["total"] == 1
    assert set(store.stats()) == {"ps1", "bat", "sh", "py", "sql"}


def test_ansi_encoded_script_is_readable(library):
    root, _, store, indexer = library
    # Windows scripts are often saved in the legacy ANSI code page, not UTF-8.
    (root / "backup.bat").write_bytes("rem sauvegarde du caf\xe9".encode("cp1252"))
    indexer.scan()
    assert store.search("sauvegarde")["total"] == 1
    assert store.document(store.search("sauvegarde")["items"][0]["id"])["content"].endswith("café")


def test_extra_text_extensions_extend_and_protect_parsers(tmp_path):
    config = tmp_path / "settings.toml"
    root = tmp_path / "documents"
    root.mkdir()
    config.write_text(
        f"folders = {json.dumps([str(root)])}\n"
        f"data_dir = {json.dumps(str(tmp_path / 'index'))}\n"
        'extra_text_extensions = ["GO", ".java"]\n'
    )
    settings = load_settings(config)
    assert settings.extra_text_extensions == (".go", ".java")
    assert {".go", ".java", ".ps1", ".py"} <= settings.text_extensions
    assert not settings.text_extensions & {".pdf", ".docx"}
    (root / "main.go").write_text("package main // deployment helper")
    store = Store(settings.data_dir)
    assert Indexer(settings, store).scan()["indexed"] == 1
    assert store.search("deployment")["total"] == 1
    for text in ('extra_text_extensions = ".py"', 'extra_text_extensions = [".docx"]'):
        config.write_text(f"folders = {json.dumps([str(root)])}\n{text}\n")
        with pytest.raises(ValueError):
            load_settings(config)


def test_office_and_html_extraction(tmp_path):
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches

    doc = Document()
    doc.add_paragraph("Paragraph sample")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Table sample"
    path = tmp_path / "example.docx"
    doc.save(path)
    assert "Paragraph sample" in extract_text(path, 1000)
    assert "Table sample" in extract_text(path, 1000)

    book = Workbook()
    book.active.append(["Invoice", 42])
    path = tmp_path / "example.xlsx"
    book.save(path)
    book.close()
    assert "Invoice\t42" in extract_text(path, 1000)

    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "Roadmap slide"
    path = tmp_path / "example.pptx"
    slides.save(path)
    assert "Roadmap slide" in extract_text(path, 1000)

    path = tmp_path / "example.html"
    path.write_text(
        '<p>Visible &amp; useful</p><script>alert("hidden")</script><style>hidden{}</style>'
    )
    text = extract_text(path, 1000)
    assert "Visible & useful" in text
    assert "hidden" not in text


def test_pdf_and_truncation(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 50 700 Td (Searchable PDF example) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    path = tmp_path / "example.pdf"
    writer.write(path)
    assert "Searchable PDF example" in extract_text(path, 1000)
    writer.encrypt("password")
    writer.write(path)
    with pytest.raises(ValueError, match="Encrypted"):
        extract_text(path, 1000)
    path = tmp_path / "text.txt"
    path.write_text("a" * 100)
    assert len(extract_text(path, 20)) == 20
    path.write_bytes("Unicode café".encode("utf-16"))
    assert extract_text(path, 100) == "Unicode café"

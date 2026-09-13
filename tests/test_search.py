import json
import os
import re
import sqlite3
import struct
import subprocess
import sys
from datetime import datetime, timezone

import pytest
from fastapi.testclient import TestClient
from outlook_msg import build as build_msg

from xtoo import mail, mcp_server
from xtoo.config import Settings, load_settings
from xtoo.enrich import entities, thread_key
from xtoo.extract import extract_text
from xtoo.indexer import Indexer
from xtoo.migrate import enrich_documents
from xtoo.store import Store
from xtoo.web import create_app


@pytest.fixture
def library(tmp_path):
    documents = tmp_path / "OneDrive - Example Company"
    documents.mkdir()
    settings = Settings(folders=(documents,), data_dir=tmp_path / "index")
    store = Store(settings.data_dir)
    return documents, settings, store, Indexer(settings, store)


def test_search_updates_deletes_and_prefixes(library):
    root, _, store, indexer = library
    file = root / "Quarterly planning.txt"
    file.write_text("The migration schedule is approved.")
    assert indexer.scan()["indexed"] == 1
    assert store.search("migr sched")["total"] == 1
    assert store.search("quarterly")["total"] == 1
    assert store.search("migration missing")["total"] == 0
    assert store.search("migration", "pdf")["total"] == 0
    assert indexer.scan()["unchanged"] == 1
    file.write_text("Completely different revised budget.")
    assert indexer.scan()["indexed"] == 1
    assert store.search("migration")["total"] == 0
    assert store.search("budget")["total"] == 1
    file.unlink()
    assert indexer.scan()["removed"] == 1
    assert store.search()["total"] == 0


def test_literal_queries_and_pagination(library):
    root, _, store, indexer = library
    for i in range(5):
        (root / f"document-{i}.txt").write_text("OR project café")
    indexer.scan()
    assert store.search('"project" OR (cafe)*')["total"] == 5
    assert store.search("'")["total"] == 0
    assert store.search("' DROP TABLE documents; --")["total"] == 0
    page1 = store.search("project", limit=2)
    page2 = store.search("project", offset=2, limit=2)
    assert page1["total"] == 5
    assert len(page2["items"]) == 2
    assert not {d["id"] for d in page1["items"]} & {d["id"] for d in page2["items"]}


def test_exclusions_symlinks_and_size_limit(library):
    root, settings, store, indexer = library
    (root / "AppData").mkdir()
    (root / "AppData/private.txt").write_text("excluded")
    outside = root.parent / "outside.txt"
    outside.write_text("outside")
    (root / "link.txt").symlink_to(outside)
    (root / "linked-dir").symlink_to(root.parent, target_is_directory=True)
    (root / "~$temporary.docx").write_text("ignore office lock")
    large = root / "large.txt"
    large.write_text("previously indexed")
    indexer.scan()
    assert store.search()["total"] == 1
    with large.open("wb") as handle:
        handle.truncate(settings.max_file_mb * 1024 * 1024 + 1)
    result = indexer.scan()
    assert result["skipped"] == 1
    assert store.search()["total"] == 0


def test_missing_root_keeps_cache_and_removed_source_purges(library):
    root, settings, store, indexer = library
    (root / "note.txt").write_text("cached")
    indexer.scan()
    root.rename(root.with_name("temporarily-unmounted"))
    assert indexer.scan()["error_count"] == 1
    assert store.search("cached")["total"] == 1
    create_app(Settings(folders=(), data_dir=settings.data_dir), background=False)
    assert store.search()["total"] == 0


def test_traversal_error_does_not_purge_existing_files(library, monkeypatch):
    root, _, store, indexer = library
    (root / "note.txt").write_text("keep")
    indexer.scan()

    def failed_walk(path, onerror, followlinks):
        onerror(PermissionError(13, "Permission denied", str(path)))
        return iter(())

    monkeypatch.setattr("xtoo.indexer.os.walk", failed_walk)
    assert indexer.scan()["error_count"] == 1
    assert store.search("keep")["total"] == 1


def test_failed_extraction_discards_old_content_and_recovers(library, monkeypatch):
    root, _, store, indexer = library
    path = root / "file.txt"
    path.write_text("old text")
    indexer.scan()
    path.write_text("new content after edit")

    def fail(*args):
        raise PermissionError("File temporarily unreadable")

    with monkeypatch.context() as context:
        context.setattr("xtoo.indexer.extract_text", fail)
        assert indexer.scan()["error_count"] == 1
        assert store.search()["total"] == 0
    assert indexer.scan()["indexed"] == 1
    assert store.search("new")["total"] == 1


def test_scripts_are_indexed_as_text(library):
    root, _, store, indexer = library
    scripts = {
        "Deploy-Cluster.ps1": "Write-Host 'restart the broker service'",
        "cleanup.bat": "@echo off\nrem purge the broker cache",
        "rotate-logs.sh": "#!/bin/bash\n# archive the broker logs",
        "report.py": "def main():\n    print('broker summary')",
        "schema.sql": "SELECT * FROM broker_events;",
    }
    for name, text in scripts.items():
        (root / name).write_text(text)
    (root / "tool.exe").write_bytes(b"MZ binary payload")
    assert indexer.scan()["indexed"] == len(scripts)
    assert store.search("broker")["total"] == len(scripts)
    assert store.search("Write-Host")["total"] == 1
    assert store.search("restart", "ps1")["total"] == 1
    assert set(store.stats()) == {"ps1", "bat", "sh", "py", "sql"}


def test_ansi_encoded_script_is_readable(library):
    root, _, store, indexer = library
    # Windows scripts are often saved in the legacy ANSI code page, not UTF-8.
    (root / "backup.bat").write_bytes("rem sauvegarde du caf\xe9".encode("cp1252"))
    indexer.scan()
    assert store.search("sauvegarde")["total"] == 1
    assert store.document(store.search("sauvegarde")["items"][0]["id"])["content"].endswith("café")


def test_extra_text_extensions_extend_and_protect_parsers(tmp_path):
    config = tmp_path / "settings.toml"
    root = tmp_path / "documents"
    root.mkdir()
    config.write_text(
        f"folders = {json.dumps([str(root)])}\n"
        f"data_dir = {json.dumps(str(tmp_path / 'index'))}\n"
        'extra_text_extensions = ["GO", ".java"]\n'
    )
    settings = load_settings(config)
    assert settings.extra_text_extensions == (".go", ".java")
    assert {".go", ".java", ".ps1", ".py"} <= settings.text_extensions
    assert not settings.text_extensions & {".pdf", ".docx"}
    (root / "main.go").write_text("package main // deployment helper")
    store = Store(settings.data_dir)
    assert Indexer(settings, store).scan()["indexed"] == 1
    assert store.search("deployment")["total"] == 1
    for text in ('extra_text_extensions = ".py"', 'extra_text_extensions = [".docx"]'):
        config.write_text(f"folders = {json.dumps([str(root)])}\n{text}\n")
        with pytest.raises(ValueError):
            load_settings(config)


def utf16(value):
    return value.encode("utf-16-le")


def test_outlook_msg_extraction_and_search(library):
    root, _, store, indexer = library
    # PR_MESSAGE_DELIVERY_TIME as a MAPI FILETIME property.
    properties = b"\x00" * 32 + struct.pack("<IIQ", 0x0E060040, 6, 133000000000000000)
    (root / "Upgrade window_20260909_101402.msg").write_bytes(
        build_msg(
            {
                "__substg1.0_0037001F": utf16("Cluster upgrade window"),
                "__substg1.0_0C1A001F": utf16("Jane Doe"),
                "__substg1.0_5D01001F": utf16("jane.doe@example.com"),
                "__substg1.0_0E04001F": utf16("Sam Lee; Platform Team"),
                "__substg1.0_1013001F": utf16("<p>Approve the window</p><script>bad()</script>"),
                "__properties_version1.0": properties,
                "__attach_version1.0_#00000000": {"__substg1.0_3707001F": utf16("runbook.pdf")},
            }
        )
    )
    assert indexer.scan()["indexed"] == 1
    content = store.document(store.search("upgrade")["items"][0]["id"])["content"]
    assert "Subject: Cluster upgrade window" in content
    assert "From: Jane Doe <jane.doe@example.com>" in content
    assert "Date: 2022-06-18" in content  # FILETIME is decoded, not left as raw bytes
    assert "Attachments: runbook.pdf" in content
    assert "Approve the window" in content and "bad()" not in content
    for query in ("jane.doe", "platform", "runbook", "approve"):
        assert store.search(query)["total"] == 1, query
    assert store.search("cluster", "msg")["total"] == 1


def test_eml_extraction_and_unreadable_message(library):
    root, _, store, indexer = library
    (root / "reply.eml").write_bytes(
        b"From: Bob Smith <bob@example.com>\r\n"
        b"To: dave@example.com\r\n"
        b"Subject: Re: cluster upgrade window\r\n"
        b"Date: Wed, 9 Sep 2026 10:14:02 +0100\r\n"
        b'Content-Type: multipart/mixed; boundary="b1"\r\n\r\n'
        b"--b1\r\n"
        b'Content-Type: text/html; charset="utf-8"\r\n\r\n'
        b"<p>Approved &amp; scheduled.</p><style>hidden{}</style>\r\n"
        b"--b1\r\n"
        b'Content-Disposition: attachment; filename="change-request.pdf"\r\n\r\n'
        b"ABC\r\n--b1--\r\n"
    )
    (root / "truncated.msg").write_bytes(b"not a compound file")
    result = indexer.scan()
    assert result["indexed"] == 1
    assert result["error_count"] == 1  # the unreadable message is reported, not indexed
    content = store.document(store.search("scheduled")["items"][0]["id"])["content"]
    assert "Subject: Re: cluster upgrade window" in content
    assert "Attachments: change-request.pdf" in content
    assert "Approved & scheduled." in content and "hidden" not in content
    assert store.search("bob@example.com")["total"] == 1
    assert store.search("", "eml")["total"] == 1


def test_received_dates_report_export_coverage(tmp_path):
    properties = b"\x00" * 32 + struct.pack("<IIQ", 0x0E060040, 6, 133000000000000000)
    message = tmp_path / "export.msg"
    message.write_bytes(build_msg({"__properties_version1.0": properties}))
    assert mail.received(message).strftime("%Y-%m-%d") == "2022-06-18"

    reply = tmp_path / "reply.eml"
    reply.write_bytes(b"Subject: hi\r\nDate: Wed, 9 Sep 2026 10:14:02 +0100\r\n\r\nbody\r\n")
    assert mail.received(reply).strftime("%Y-%m-%d") == "2026-09-09"

    undated = tmp_path / "undated.eml"
    undated.write_bytes(b"Subject: hi\r\n\r\nbody\r\n")
    assert mail.received(undated) is None
    assert mail.received(tmp_path / "export.msg") is not None


def conversation(root, name, subject, body, when, to="Ford, Alex; Lee, Sam"):
    epoch = datetime(1601, 1, 1, tzinfo=timezone.utc)
    ticks = int((when - epoch).total_seconds() * 10_000_000)
    (root / name).write_bytes(
        build_msg(
            {
                "__substg1.0_0037001F": utf16(subject),
                "__substg1.0_0C1A001F": utf16("Jane Doe"),
                "__substg1.0_5D01001F": utf16("jane.doe@example.com"),
                "__substg1.0_0E04001F": utf16(to),
                "__substg1.0_1000001F": utf16(body),
                "__properties_version1.0": b"\x00" * 32 + struct.pack("<IIQ", 0x0E060040, 6, ticks),
            }
        )
    )


@pytest.fixture
def correspondence(library):
    root, settings, store, indexer = library
    moment = datetime(2026, 3, 2, 9, 0, tzinfo=timezone.utc)
    conversation(root, "a.msg", "[JIRA] (PROJ-4821) Upgrade fails", "First report", moment)
    conversation(
        root,
        "b.msg",
        "RE: [JIRA] (PROJ-4821) Upgrade fails",
        "Root cause",
        moment.replace(month=4),
    )
    conversation(
        root,
        "c.msg",
        "FW: RE: [JIRA] (PROJ-4821) Upgrade fails",
        "Over to you",
        moment.replace(month=5),
    )
    conversation(root, "d.msg", "Lab notice", "UTF-8 and SHA-1 are not tickets", moment)
    (root / "fix.sh").write_text("#!/bin/bash\n# workaround for PROJ-4821\nrestart vpxd\n")
    indexer.scan()
    return root, settings, store, indexer


def test_message_dates_replace_file_timestamps(correspondence):
    _, _, store, _ = correspondence
    dates = {
        item["title"]: datetime.fromtimestamp(item["document_ns"] / 1e9, timezone.utc).date()
        for item in store.search()["items"]
    }
    assert str(dates["a.msg"]) == "2026-03-02"
    assert str(dates["c.msg"]) == "2026-05-02"
    # A file with no message date keeps its filesystem timestamp.
    assert dates["fix.sh"] >= datetime(2025, 1, 1).date()
    assert [item["title"] for item in store.search()["items"]][:2] == ["fix.sh", "c.msg"]


def test_replies_collapse_into_one_conversation(correspondence):
    _, _, store, _ = correspondence
    assert store.search("upgrade")["total"] == 3
    collapsed = store.search("upgrade", collapse=True)
    assert collapsed["total"] == 1
    assert collapsed["items"][0]["thread_size"] == 3
    # Documents outside a conversation are never merged together.
    assert store.search(collapse=True)["total"] == 3
    assert thread_key("Subject: RE: FW: Budget\n", "x.msg", "msg") == "budget"
    assert thread_key("Subject: Hi\n", "x.msg", "txt") == ""


def test_entities_link_mail_and_scripts(correspondence):
    _, _, store, _ = correspondence
    linked = store.search(entity="PROJ-4821")
    assert {item["title"] for item in linked["items"]} == {"a.msg", "b.msg", "c.msg", "fix.sh"}
    assert store.search(entity="Lee, Sam")["total"] == 4
    assert store.search(entity="jane.doe@example.com")["total"] == 4
    assert store.entities("PROJ")[0] == {"kind": "ticket", "name": "PROJ-4821", "count": 4}
    assert store.entities("%") == []  # wildcards are escaped, not interpreted
    found = {name for _, name in entities("UTF-8 SHA-1 PDF-2 see OPS-77", "txt")}
    assert found == {"OPS-77"}


def test_mcp_tools_search_read_and_link(correspondence):
    _, _, store, _ = correspondence
    found = mcp_server.search(store, "upgrade", limit=5)
    assert found["total"] == 1 and found["results"][0]["messages_in_conversation"] == 3
    assert found["results"][0]["date"] == "2026-03-02"
    document = mcp_server.read(store, found["results"][0]["id"], max_characters=200)
    assert "ticket:PROJ-4821" in document["entities"] and document["truncated"] is False
    assert mcp_server.by_entity(store, "PROJ-4821")["total"] == 4
    assert "ticket:PROJ-4821 (4)" in mcp_server.names(store, "PROJ")["names"]
    assert mcp_server.read(store, 99999)["error"]


def test_migrate_enriches_an_index_built_before_enrichment(tmp_path):
    directory = tmp_path / "index"
    directory.mkdir()
    legacy = sqlite3.connect(directory / "index.sqlite3")
    legacy.executescript(
        """CREATE TABLE documents (id INTEGER PRIMARY KEY, path TEXT NOT NULL UNIQUE,
        root TEXT NOT NULL, title TEXT NOT NULL, kind TEXT NOT NULL, modified_ns INTEGER NOT NULL,
        size INTEGER NOT NULL, content TEXT NOT NULL);
        CREATE VIRTUAL TABLE search_index USING fts5(title, content, content='documents',
        content_rowid='id');
        CREATE TRIGGER documents_insert AFTER INSERT ON documents BEGIN
        INSERT INTO search_index(rowid, title, content) VALUES (new.id, new.title, new.content);
        END;"""
    )
    legacy.execute(
        "INSERT INTO documents(path, root, title, kind, modified_ns, size, content) "
        "VALUES ('/m/one.msg', '/m', 'one.msg', 'msg', 1700000000000000000, 10, ?)",
        (
            "Subject: RE: Cluster upgrade\nFrom: Jane Doe <jane.doe@example.com>\n"
            "To: Lee, Sam\nDate: 2026-03-02 09:00 UTC\nSee PROJ-4821 for detail.",
        ),
    )
    legacy.commit()
    legacy.close()

    store = Store(directory)  # opening an old index adds the new columns
    before = store.search("cluster")["items"][0]
    assert before["document_ns"] == 1700000000000000000  # falls back to the file timestamp
    assert enrich_documents(store) == 1
    after = store.search("cluster")["items"][0]
    assert datetime.fromtimestamp(after["document_ns"] / 1e9, timezone.utc).date().isoformat() == (
        "2026-03-02"
    )
    assert after["thread"] == "cluster upgrade"
    assert store.search(entity="PROJ-4821")["total"] == 1
    assert enrich_documents(store) == 0  # nothing left to do on a second run


def test_api_exposes_conversations_and_entity_links(correspondence):
    _, settings, _, _ = correspondence
    with TestClient(create_app(settings, background=False), base_url="http://localhost") as client:
        grouped = client.get("/api/search", params={"q": "upgrade", "collapse": "true"}).json()
        assert grouped["total"] == 1
        assert grouped["items"][0]["thread_size"] == 3
        assert client.get("/api/search", params={"q": "upgrade"}).json()["total"] == 3
        linked = client.get("/api/search", params={"entity": "PROJ-4821"}).json()
        assert linked["total"] == 4
        names = client.get("/api/entities", params={"prefix": "PROJ"}).json()["items"]
        assert names[0] == {"kind": "ticket", "name": "PROJ-4821", "count": 4}
        assert client.get("/api/entities").json()["items"]  # no prefix browses the busiest
        document = client.get(f"/api/documents/{linked['items'][0]['id']}").json()
        assert {"kind": "ticket", "name": "PROJ-4821"} in document["entities"]


def bag_of_words(texts):
    """A deterministic stand-in for the embedding model: no download, no network."""
    encoded = []
    for text in texts:
        vector = [0.0] * 256
        for word in set(re.findall(r"\w+", text.lower())):
            vector[hash(word) % 256] += 1.0
        scale = sum(value * value for value in vector) ** 0.5 or 1.0
        encoded.append([value / scale for value in vector])
    return encoded


def test_semantic_vectors_build_resume_and_rank(correspondence):
    pytest.importorskip("sqlite_vec")
    from xtoo import vectors

    root, _, store, indexer = correspondence
    built = vectors.build(store, encode=bag_of_words)
    assert built["documents"] == 5 and built["chunks"] >= 5
    assert vectors.build(store, encode=bag_of_words)["documents"] == 0  # resumes, does not redo

    found = vectors.similar(store, "restart vpxd", encode=bag_of_words)
    script = store.search("workaround")["items"][0]["id"]
    assert script in found

    fused = vectors.search(store, "vpxd", limit=5, encode=bag_of_words)
    assert fused["total"] >= 1
    assert all(item["snippet"] for item in fused["items"])

    # Editing a file makes its vectors stale, and only that document is redone.
    (root / "fix.sh").write_text("#!/bin/bash\n# replaced entirely\n")
    indexer.scan()
    assert vectors.build(store, encode=bag_of_words)["documents"] == 1

    # Deleting a file clears its vectors rather than leaving them to be matched.
    (root / "fix.sh").unlink()
    indexer.scan()
    vectors.build(store, encode=bag_of_words)
    assert script not in vectors.similar(store, "restart vpxd", encode=bag_of_words)


def test_chunking_covers_the_start_of_a_long_document():
    pytest.importorskip("sqlite_vec")
    from xtoo import vectors

    assert vectors.pieces("title", "") == ["title"]
    windows = vectors.pieces("report", "word " * 4000)
    assert len(windows) == vectors.MAX_CHUNKS
    assert all(len(window) <= vectors.CHUNK for window in windows)
    assert windows[0].startswith("report")
    # Windows overlap, so a phrase on a boundary is not lost.
    assert windows[0][-vectors.OVERLAP :] in windows[1]


def test_configuration_validation_and_nested_roots(tmp_path):
    config = tmp_path / "settings.toml"
    root = tmp_path / "documents"
    config.write_text(f"folders = {json.dumps([str(root), str(root / 'nested'), str(root)])}\n")
    assert load_settings(config).folders == (root,)
    for text in (
        "folders = []",
        "folders = [1]",
        'folders = ["/tmp"]\ninterval_seconds = 0',
        'folders = ["/tmp"]\nmisspelled = 1',
    ):
        config.write_text(text)
        with pytest.raises(ValueError):
            load_settings(config)


def test_office_and_html_extraction(tmp_path):
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches

    doc = Document()
    doc.add_paragraph("Paragraph sample")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Table sample"
    path = tmp_path / "example.docx"
    doc.save(path)
    assert "Paragraph sample" in extract_text(path, 1000)
    assert "Table sample" in extract_text(path, 1000)

    book = Workbook()
    book.active.append(["Invoice", 42])
    path = tmp_path / "example.xlsx"
    book.save(path)
    book.close()
    assert "Invoice\t42" in extract_text(path, 1000)

    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "Roadmap slide"
    path = tmp_path / "example.pptx"
    slides.save(path)
    assert "Roadmap slide" in extract_text(path, 1000)

    path = tmp_path / "example.html"
    path.write_text(
        '<p>Visible &amp; useful</p><script>alert("hidden")</script><style>hidden{}</style>'
    )
    text = extract_text(path, 1000)
    assert "Visible & useful" in text
    assert "hidden" not in text


def test_pdf_and_truncation(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 50 700 Td (Searchable PDF example) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    path = tmp_path / "example.pdf"
    writer.write(path)
    assert "Searchable PDF example" in extract_text(path, 1000)
    writer.encrypt("password")
    writer.write(path)
    with pytest.raises(ValueError, match="Encrypted"):
        extract_text(path, 1000)
    path = tmp_path / "text.txt"
    path.write_text("a" * 100)
    assert len(extract_text(path, 20)) == 20
    path.write_bytes("Unicode café".encode("utf-16"))
    assert extract_text(path, 100) == "Unicode café"


def test_api_search_preview_and_request_guards(library):
    root, settings, _, indexer = library
    (root / "report.txt").write_text('<script>alert("example")</script> budget')
    indexer.scan()
    with TestClient(create_app(settings, background=False), base_url="http://localhost") as client:
        response = client.get("/")
        assert response.status_code == 200
        assert "frame-ancestors 'none'" in response.headers["content-security-policy"]
        assert client.get("/static/app.js").status_code == 200
        result = client.get("/api/search", params={"q": "budg"}).json()
        assert result["total"] == 1
        doc = client.get(f"/api/documents/{result['items'][0]['id']}").json()
        assert "<script>" in doc["content"]  # JSON; UI uses text nodes, never HTML.
        assert not doc["preview_truncated"]
        assert client.get("/api/documents/9999").status_code == 404
        assert client.get("/api/search?offset=-1").status_code == 422
        assert client.get("/api/search?limit=1000").status_code == 422
        assert client.get("/api/status").json()["total_documents"] == 1
        assert client.get("/api/search", headers={"Host": "attacker.example"}).status_code == 400
        assert client.post("/api/index").status_code == 403
        headers = {"X-Xtoo-Request": "1", "Origin": "https://attacker.example"}
        assert client.post("/api/index", headers=headers).status_code == 403
        headers["Origin"] = "http://localhost"
        assert client.post("/api/index", headers=headers).status_code == 202
        preflight = client.options(
            "/api/index",
            headers={"Origin": "https://attacker.example", "Access-Control-Request-Method": "POST"},
        )
        assert "access-control-allow-origin" not in preflight.headers


def test_background_scan_starts_and_stops(library):
    import time

    root, settings, _, _ = library
    (root / "note.txt").write_text("background")
    app = create_app(settings)
    with TestClient(app, base_url="http://localhost") as client:
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            if client.get("/api/status").json()["last_finished"]:
                break
            time.sleep(0.02)
        assert client.get("/api/search?q=background").json()["total"] == 1
        (root / "note.txt").unlink()
        client.post("/api/index", headers={"X-Xtoo-Request": "1"})
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            if client.get("/api/search").json()["total"] == 0:
                break
            time.sleep(0.02)
        assert client.get("/api/search").json()["total"] == 0
    assert not app.state.indexer._thread.is_alive()


def test_cli_init_keeps_configuration_outside_checkout(tmp_path):
    folder = tmp_path / "Documents with spaces"
    folder.mkdir()
    env = {
        **os.environ,
        "XDG_CONFIG_HOME": str(tmp_path / "config"),
        "XDG_DATA_HOME": str(tmp_path / "data"),
    }
    command = [sys.executable, "-m", "xtoo.cli", "init", "--folder", str(folder)]
    result = subprocess.run(command, env=env, capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    config = tmp_path / "config/xtoo/config.toml"
    assert load_settings(config).folders == (folder,)
    assert config.stat().st_mode & 0o077 == 0
    assert subprocess.run(command, env=env, capture_output=True).returncode == 2
    result = subprocess.run(
        [sys.executable, "-m", "xtoo.cli", "index"], env=env, capture_output=True, text=True
    )
    assert result.returncode == 0, result.stderr
    assert json.loads(result.stdout)["error_count"] == 0

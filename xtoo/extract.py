"""Text extraction only: never execute macros, scripts, or embedded objects."""

from html.parser import HTMLParser
from pathlib import Path
from zipfile import ZipFile

TEXT_EXTENSIONS = {".txt", ".md", ".rst", ".csv", ".tsv", ".json", ".xml", ".yaml", ".yml", ".log"}
SUPPORTED = TEXT_EXTENSIONS | {".html", ".htm", ".pdf", ".docx", ".xlsx", ".pptx"}


class HTMLText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style"}:
            self.hidden += 1
        elif tag in {"p", "div", "br", "li", "h1", "h2", "tr"}:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style"}:
            self.hidden = max(0, self.hidden - 1)

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def chunks(path: Path):
    suffix = path.suffix.lower()
    if suffix in TEXT_EXTENSIONS | {".html", ".htm"}:
        data = path.read_bytes()
        text = data.decode(
            "utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig",
            errors="replace",
        )
        if suffix in {".html", ".htm"}:
            parser = HTMLText()
            parser.feed(text)
            text = "".join(parser.parts)
        yield text
    elif suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("Encrypted PDF; unlock a permitted copy before indexing")
        for page in reader.pages:
            yield page.extract_text() or ""
    else:
        # Reject unusually large expanded Office archives before invoking parsers.
        with ZipFile(path) as archive:
            if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Expanded Office document exceeds 100 MB")
        if suffix == ".docx":
            from docx import Document

            document = Document(path)
            for paragraph in document.paragraphs:
                yield paragraph.text
            for table in document.tables:
                for row in table.rows:
                    yield "\t".join(cell.text for cell in row.cells)
        elif suffix == ".xlsx":
            from openpyxl import load_workbook

            book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            try:
                for sheet in book:
                    yield sheet.title
                    for row in sheet.iter_rows(values_only=True):
                        yield "\t".join(str(value) if value is not None else "" for value in row)
            finally:
                book.close()
        elif suffix == ".pptx":
            from pptx import Presentation

            for number, slide in enumerate(Presentation(path).slides, 1):
                yield f"Slide {number}"
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        yield shape.text
                    if shape.has_table:
                        for row in shape.table.rows:
                            yield "\t".join(cell.text for cell in row.cells)
        else:
            raise ValueError(f"Unsupported format: {suffix}")


def extract_text(path: Path, limit: int) -> str:
    parts = []
    remaining = limit
    iterator = chunks(path)
    try:
        for part in iterator:
            part = part.replace("\x00", "")[:remaining]
            parts.append(part)
            remaining -= len(part) + 1
            if remaining <= 0:
                break
    finally:
        iterator.close()
    return "\n".join(parts)[:limit]
